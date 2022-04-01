import os
import re
from openpyxl import Workbook
from openpyxl import load_workbook

def data_to_dict(file_path):


	dict_var={}
	dev_list=[]
	#print("Hello")
	for i in os.listdir(file_path):
		#print(i)
		if i:
			count=1
			print(i)
			commd=[]
		
		#commd_out[count]={}
			wb2=load_workbook('/Users/praette/Desktop/Project_python/{}'.format(i))
			sheet1=wb2['Sheet1']
			row_len=len(sheet1['B'])
			col_len=len(sheet1['1'])
			row_1=sheet1['1']
			col_b=sheet1['B']
			#print(row_len)
			#print(col_len)
		
			for x in range(1,row_len):
				dev_list.append(sheet1['B'][x].value)
			for y in range(1,col_len):
				commd.append(row_1[y].value)
			for dev in dev_list:
				dict_var[dev]={}
			#print(dict_var.keys())
			for z in range(2,row_len):
				row_num=z
				col_count=0
				row_x=sheet1[z+1]
			

				for i in commd:
					col_count=col_count+1
					dict_var[col_b[z].value][i]=row_x[col_count].value
	return dict_var			

		##print(str_out)
def device_versions(dict_var):
	versions=Workbook()
	ws=versions.create_sheet("version")
	row_count=1
	for i  in dict_var.keys():
			
		ws.cell(row=row_count,column=1).value=i
		str_out=dict_var[i].get('show version')
		str_out=str(str_out)
		#print(type(str_out))
		str_out=str_out.split('\n')[0]
		if re.search(r'IOS XE',str_out):
			pattern=re.findall(r'\d+\.\d+\.\d+\w?\.\w',str_out)
			if len(pattern)!=0:
				ws.cell(row_count,2).value=pattern[0]
			else:
				ws.cell(row_count,2).value="None"

		row_count=row_count+1


	versions.save('new_big_file.xlsx')

			
	return versions

def WAN_resiliency:
import re
from openpyxl import Workbook
from openpyxl import load_workbook


string ='''BGP table version is 5, local router ID is 44.44.44.44
Status codes: s suppressed, d damped, h history, * valid, > best, i - internal, 
              r RIB-failure, S Stale, m multipath, b backup-path, f RT-Filter, 
              x best-external, a additional-path, c RIB-compressed, 
Origin codes: i - IGP, e - EGP, ? - incomplete
RPKI validation codes: V valid, I invalid, N Not found

     Network          Next Hop            Metric LocPrf Weight Path
 *>i 1.1.1.1/32       192.168.3.3              0    100      0 i
 *>i 2.2.2.2/32       192.168.3.3              0    100      0 i

Total number of prefixes 2'''
#str1=str(string)

file_name="/Users/praette/Desktop/WAN_res.xlsx"
wb2=load_workbook("/Users/praette/Desktop/WAN_res.xlsx")
sheet1=wb2['Sheet1']
str1=sheet1['A'][0].value
detials=str1.split('\n')
print(detials)
x=[]
for i in detials:
	pattern=re.search("\*>i ([\d+\.]+\/\d+)",i)
	if pattern:
		x.append(pattern.group(0))

print(x)


y=['2.2.2.2/32','3.3.3.3/32']
x=set(x)
y=set(y)
print(x-y)
		
			
file_path="/Users/praette/Desktop/Project_python"
wb1=Workbook()
dict_obj=data_to_dict(file_path)
wb1=device_versions(dict_obj)


